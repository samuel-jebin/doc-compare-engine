from docx import Document


def read_from_file(file_path):
    doc = Document(file_path)

    for para in doc.paragraphs:
        text = "\n".join([para.text for para in doc.paragraphs])
    
    print(text)
    return text

# file_path = r"C:\Users\K1194\Downloads\test-1.docx"
# read_from_file(file_path)